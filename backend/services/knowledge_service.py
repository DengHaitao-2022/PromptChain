"""知识库解析、索引与检索服务。"""

from __future__ import annotations

import hashlib
import logging
import math
import re
import uuid
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from sqlalchemy import and_, delete, or_, select, update
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from core.config import get_settings
from core.time import utc_now_naive
from models.auth_models import MemberRole
from models.knowledge import (
    EvidenceChunk,
    EvidencePack,
    KnowledgeBase,
    KnowledgeBaseStatus,
    KnowledgeChunk,
    KnowledgeConflict,
    KnowledgeDocument,
    KnowledgeDocumentLifecycleStatus,
    KnowledgeDocumentStatus,
    KnowledgeScope,
    KnowledgeSearchRequest,
    KnowledgeSearchResponse,
    KnowledgeUsageStats,
    RetrievalEvaluationRequest,
    RetrievalEvaluationResponse,
    RetrievalEvaluationResult,
    RetrievalEvaluationSummary,
    RetrievalMode,
)
from orm.knowledge_orm import (
    KnowledgeBaseORM,
    KnowledgeChunkORM,
    KnowledgeDocumentORM,
    KnowledgeEmbeddingORM,
    KnowledgeRetrievalLogORM,
)
from services.llm_usage import estimate_tokens
from services.permission_service import check_permission

TOKEN_PATTERN = re.compile(r"[\w\u4e00-\u9fff]+", re.UNICODE)
SUPPORTED_TEXT_TYPES = {"txt", "md", "markdown", "text"}
SUPPORTED_DOCUMENT_TYPES = {*SUPPORTED_TEXT_TYPES, "docx", "pdf"}
DOCX_ZIP_MAGIC = b"PK"
PDF_MAGIC = b"%PDF"
logger = logging.getLogger(__name__)


def _normalize_file_type(file_name: str) -> str:
    suffix = Path(file_name).suffix.lower().lstrip(".")
    if suffix == "markdown":
        return "md"
    return suffix or "txt"


def _validate_upload_content(file_type: str, content: bytes) -> None:
    """在解析前做轻量上传安全检查，避免明显伪造文件进入索引流程。"""
    settings = get_settings()
    if not content:
        raise ValueError("文档内容不能为空")
    if len(content) > settings.KNOWLEDGE_MAX_UPLOAD_BYTES:
        max_mb = settings.KNOWLEDGE_MAX_UPLOAD_BYTES / 1024 / 1024
        raise ValueError(f"文档大小不能超过 {max_mb:.0f} MB")
    if file_type == "pdf" and not content.startswith(PDF_MAGIC):
        raise ValueError("PDF 文件格式校验失败")
    if file_type == "docx" and not content.startswith(DOCX_ZIP_MAGIC):
        raise ValueError("DOCX 文件格式校验失败")


def validate_upload_file(file_name: str, content: bytes) -> str:
    """统一校验知识库上传文件，并返回归一化后的文件类型。"""
    file_type = _normalize_file_type(file_name)
    if file_type not in SUPPORTED_DOCUMENT_TYPES:
        raise ValueError(f"暂不支持的文档类型: {file_type}")
    _validate_upload_content(file_type, content)
    return file_type


def _hash_content(content: bytes | str) -> str:
    raw = content.encode("utf-8") if isinstance(content, str) else content
    return hashlib.sha256(raw).hexdigest()


def _tokenize(text: str) -> list[str]:
    return [token.lower() for token in TOKEN_PATTERN.findall(text)]


def _cosine_similarity(left: list[float], right: list[float]) -> float:
    if not left or not right or len(left) != len(right):
        return 0.0
    dot = sum(a * b for a, b in zip(left, right, strict=False))
    left_norm = math.sqrt(sum(value * value for value in left))
    right_norm = math.sqrt(sum(value * value for value in right))
    if left_norm == 0 or right_norm == 0:
        return 0.0
    return max(0.0, min(1.0, dot / (left_norm * right_norm)))


def _keyword_score(query: str, content: str) -> float:
    query_terms = Counter(_tokenize(query))
    if not query_terms:
        return 0.0
    content_terms = Counter(_tokenize(content))
    matched = sum(min(count, content_terms.get(term, 0)) for term, count in query_terms.items())
    return min(1.0, matched / max(1, sum(query_terms.values())))


def _normalize_role_value(role: Any) -> MemberRole | str:
    """兼容枚举和测试中常见的字符串角色。"""
    try:
        return MemberRole(str(getattr(role, "value", role)))
    except ValueError:
        return str(getattr(role, "value", role))


def _is_workspace_scope(scope: KnowledgeScope) -> bool:
    """工作空间库是团队公共事实源，权限要求高于个人库。"""
    return scope == KnowledgeScope.WORKSPACE


def _row_to_knowledge_base(row: KnowledgeBaseORM) -> KnowledgeBase:
    return KnowledgeBase(
        id=row.id,
        workspace_id=row.workspace_id,
        owner_user_id=row.owner_user_id,
        workflow_run_id=getattr(row, "workflow_run_id", None),
        scope=KnowledgeScope(row.scope),
        name=row.name,
        description=row.description,
        status=KnowledgeBaseStatus(row.status),
        created_by=row.created_by,
        created_at=row.created_at,
        updated_at=row.updated_at,
    )


def _row_to_document(row: KnowledgeDocumentORM) -> KnowledgeDocument:
    return KnowledgeDocument(
        id=row.id,
        kb_id=row.kb_id,
        workspace_id=row.workspace_id,
        owner_user_id=row.owner_user_id,
        workflow_run_id=getattr(row, "workflow_run_id", None),
        file_name=row.file_name,
        file_type=row.file_type,
        storage_uri=row.storage_uri,
        checksum=row.checksum,
        version=row.version,
        status=KnowledgeDocumentLifecycleStatus(getattr(row, "status", None) or "active"),
        parse_status=KnowledgeDocumentStatus(row.parse_status),
        index_status=KnowledgeDocumentStatus(row.index_status),
        error_message=row.error_message,
        metadata=getattr(row, "metadata_json", None) or {},
        created_by=row.created_by,
        created_at=row.created_at,
        updated_at=row.updated_at,
    )


def _row_to_chunk(row: KnowledgeChunkORM) -> KnowledgeChunk:
    return KnowledgeChunk(
        id=row.id,
        document_id=row.document_id,
        kb_id=row.kb_id,
        workspace_id=row.workspace_id,
        owner_user_id=row.owner_user_id,
        workflow_run_id=getattr(row, "workflow_run_id", None),
        chunk_index=row.chunk_index,
        content=row.content,
        token_count=row.token_count,
        heading_path=row.heading_path or [],
        page_number=row.page_number,
        metadata=row.metadata_json or {},
        content_hash=row.content_hash,
        created_at=row.created_at,
    )


class DocumentParser:
    """把上传文件解析成可索引文本。"""

    def parse(self, file_name: str, content: bytes) -> str:
        file_type = _normalize_file_type(file_name)
        if file_type in SUPPORTED_TEXT_TYPES:
            return content.decode("utf-8", errors="ignore")
        if file_type == "docx":
            return self._parse_docx(content)
        if file_type == "pdf":
            return self._parse_pdf(content)
        raise ValueError(f"暂不支持的文档类型: {file_type}")

    def _parse_docx(self, content: bytes) -> str:
        from io import BytesIO

        from docx import Document

        document = Document(BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)

    def _parse_pdf(self, content: bytes) -> str:
        from io import BytesIO

        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        pages: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"\n\n[page:{index}]\n{text.strip()}")
        return "\n".join(pages)


class Chunker:
    """按标题优先、超长滑窗的方式切分文档。"""

    def __init__(self, target_tokens: int | None = None, overlap_tokens: int | None = None):
        settings = get_settings()
        self.target_tokens = target_tokens or settings.KNOWLEDGE_CHUNK_TARGET_TOKENS
        self.overlap_tokens = overlap_tokens or settings.KNOWLEDGE_CHUNK_OVERLAP_TOKENS

    def split(self, text: str) -> list[dict[str, Any]]:
        sections = self._split_by_heading(text)
        chunks: list[dict[str, Any]] = []
        for section in sections:
            chunks.extend(self._split_long_section(section))
        return chunks

    def _split_by_heading(self, text: str) -> list[dict[str, Any]]:
        sections: list[dict[str, Any]] = []
        current_heading: list[str] = []
        current_lines: list[str] = []
        current_page: int | None = None

        def flush() -> None:
            content = "\n".join(current_lines).strip()
            if content:
                sections.append(
                    {
                        "content": content,
                        "heading_path": list(current_heading),
                        "page_number": current_page,
                    }
                )

        for raw_line in text.splitlines():
            line = raw_line.strip()
            page_match = re.match(r"^\[page:(\d+)\]$", line)
            if page_match:
                current_page = int(page_match.group(1))
                continue

            heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading_match:
                flush()
                level = len(heading_match.group(1))
                title = heading_match.group(2).strip()
                current_heading = [*current_heading[: level - 1], title]
                current_lines = [title]
                continue

            current_lines.append(raw_line)

        flush()
        if not sections and text.strip():
            sections.append({"content": text.strip(), "heading_path": [], "page_number": None})
        return sections

    def _split_long_section(self, section: dict[str, Any]) -> list[dict[str, Any]]:
        content = section["content"]
        token_count = estimate_tokens(content)
        if token_count <= self.target_tokens:
            return [{**section, "token_count": token_count}]

        # estimate_tokens 约等于每 2 字符 1 token，因此滑窗按字符长度近似即可。
        window = max(200, self.target_tokens * 2)
        overlap = max(0, self.overlap_tokens * 2)
        step = max(1, window - overlap)
        chunks = []
        for start in range(0, len(content), step):
            chunk_text = content[start : start + window].strip()
            if not chunk_text:
                continue
            chunks.append(
                {
                    **section,
                    "content": chunk_text,
                    "token_count": estimate_tokens(chunk_text),
                    "metadata": {"offset_start": start, "offset_end": start + len(chunk_text)},
                }
            )
            if start + window >= len(content):
                break
        return chunks


class EmbeddingProvider:
    """Embedding provider；默认本地 hash，可通过环境变量切换 OpenAI。"""

    def __init__(self):
        self.settings = get_settings()
        self._last_provider_name = self.settings.KNOWLEDGE_EMBEDDING_PROVIDER
        self._last_model_name = self.settings.KNOWLEDGE_EMBEDDING_MODEL

    @property
    def model_name(self) -> str:
        return self.settings.KNOWLEDGE_EMBEDDING_MODEL

    @property
    def effective_provider_name(self) -> str:
        return self._last_provider_name

    @property
    def effective_model_name(self) -> str:
        return self._last_model_name

    @property
    def dimension(self) -> int:
        return self.settings.KNOWLEDGE_EMBEDDING_DIMENSION

    async def embed_query(self, text: str) -> list[float]:
        return (await self.embed_documents([text]))[0]

    async def embed_documents(self, texts: list[str]) -> list[list[float]]:
        if self.settings.KNOWLEDGE_EMBEDDING_PROVIDER == "openai":
            try:
                from langchain_openai import OpenAIEmbeddings

                embeddings = OpenAIEmbeddings(
                    model=self.model_name,
                    api_key=self.settings.OPENAI_API_KEY,
                    dimensions=self.dimension,
                )
                self._last_provider_name = "openai"
                self._last_model_name = self.model_name
                return await embeddings.aembed_documents(texts)
            except Exception as exc:
                if not self.settings.KNOWLEDGE_EMBEDDING_FALLBACK_TO_HASH:
                    raise RuntimeError(f"OpenAI embedding 调用失败: {exc}") from exc
                logger.warning("OpenAI embedding 不可用，已回退本地 hash provider: %s", exc)
                return self._embed_with_hash(texts)
        return self._embed_with_hash(texts)

    def _embed_with_hash(self, texts: list[str]) -> list[list[float]]:
        self._last_provider_name = "hash"
        self._last_model_name = "promptchain-hash-embedding-v1"
        return [self._hash_embedding(text) for text in texts]

    def _hash_embedding(self, text: str) -> list[float]:
        vector = [0.0] * self.dimension
        tokens = _tokenize(text) or [text]
        for token in tokens:
            digest = hashlib.sha256(token.encode("utf-8")).digest()
            index = int.from_bytes(digest[:4], "big") % self.dimension
            sign = 1.0 if digest[4] % 2 == 0 else -1.0
            vector[index] += sign
        norm = math.sqrt(sum(value * value for value in vector)) or 1.0
        return [round(value / norm, 8) for value in vector]


class RerankerProvider:
    """Reranker provider 抽象，MVP 默认使用确定性启发式重排。"""

    provider_name = "heuristic"

    def rerank(self, query: str, candidates: list[_Candidate]) -> list[_Candidate]:
        query_terms = set(_tokenize(query))
        for candidate in candidates:
            heading_terms = set(_tokenize(" ".join(candidate.chunk.heading_path or [])))
            title_terms = set(_tokenize(candidate.document.file_name))
            authority_boost = 0.04 if candidate.kb.scope == KnowledgeScope.WORKSPACE.value else 0.0
            heading_boost = 0.03 if query_terms & heading_terms else 0.0
            title_boost = 0.03 if query_terms & title_terms else 0.0
            rerank_score = min(1.0, candidate.score + authority_boost + heading_boost + title_boost)
            candidate.rerank_score = round(rerank_score, 4)
            candidate.score = candidate.rerank_score
        return candidates


@dataclass
class _Candidate:
    chunk: KnowledgeChunkORM
    document: KnowledgeDocumentORM
    kb: KnowledgeBaseORM
    embedding: KnowledgeEmbeddingORM | None
    score: float
    vector_score: float
    keyword_score: float
    rerank_score: float | None = None


class KnowledgeService:
    """知识库应用服务。"""

    def __init__(self, session: AsyncSession):
        self.session = session
        self.parser = DocumentParser()
        self.chunker = Chunker()
        self.embedding_provider = EmbeddingProvider()
        self.reranker_provider = RerankerProvider()

    def _ensure_kb_operation_allowed(
        self,
        kb: KnowledgeBaseORM,
        *,
        user_id: str,
        role: Any,
        action: str,
    ) -> None:
        """按 scope 做知识库操作授权，个人库 owner 不受工作空间写权限限制。"""
        role_value = _normalize_role_value(role)
        if kb.scope in {KnowledgeScope.PERSONAL.value, KnowledgeScope.RUN_UPLOAD.value}:
            if kb.owner_user_id == user_id:
                return
            raise ValueError("知识库不存在或无权访问")

        if check_permission(role_value, "knowledge_base", action):
            return
        raise ValueError(f"您没有 knowledge_base.{action} 的权限")

    async def list_knowledge_bases(
        self,
        *,
        workspace_id: str,
        user_id: str,
        scopes: list[KnowledgeScope] | None = None,
    ) -> list[KnowledgeBase]:
        filters = [KnowledgeBaseORM.workspace_id == workspace_id]
        if scopes:
            filters.append(KnowledgeBaseORM.scope.in_([scope.value for scope in scopes]))
        else:
            filters.append(KnowledgeBaseORM.scope != KnowledgeScope.RUN_UPLOAD.value)
        filters.append(
            or_(
                KnowledgeBaseORM.scope == KnowledgeScope.WORKSPACE.value,
                KnowledgeBaseORM.owner_user_id == user_id,
            )
        )
        result = await self.session.execute(
            select(KnowledgeBaseORM)
            .where(and_(*filters))
            .order_by(KnowledgeBaseORM.updated_at.desc())
        )
        return [_row_to_knowledge_base(row) for row in result.scalars().all()]

    async def create_knowledge_base(
        self,
        *,
        workspace_id: str,
        user_id: str,
        role: Any,
        name: str,
        description: str | None,
        scope: KnowledgeScope,
        workflow_run_id: str | None = None,
    ) -> KnowledgeBase:
        if scope == KnowledgeScope.RUN_UPLOAD and not workflow_run_id:
            raise ValueError("本次运行资料知识库必须绑定 workflow_run_id")
        if _is_workspace_scope(scope) and not check_permission(
            _normalize_role_value(role),
            "knowledge_base",
            "create",
        ):
            raise ValueError("您没有 knowledge_base.create 的权限")
        owner_user_id = (
            user_id if scope in {KnowledgeScope.PERSONAL, KnowledgeScope.RUN_UPLOAD} else None
        )
        row = KnowledgeBaseORM(
            id=str(uuid.uuid4()),
            workspace_id=workspace_id,
            owner_user_id=owner_user_id,
            workflow_run_id=workflow_run_id if scope == KnowledgeScope.RUN_UPLOAD else None,
            scope=scope.value,
            name=name,
            description=description,
            status=KnowledgeBaseStatus.ACTIVE.value,
            created_by=user_id,
        )
        self.session.add(row)
        await self.session.commit()
        await self.session.refresh(row)
        return _row_to_knowledge_base(row)

    async def get_knowledge_base(
        self,
        *,
        kb_id: str,
        workspace_id: str,
        user_id: str,
    ) -> KnowledgeBaseORM:
        result = await self.session.execute(
            select(KnowledgeBaseORM).where(
                KnowledgeBaseORM.id == kb_id,
                KnowledgeBaseORM.workspace_id == workspace_id,
                or_(
                    KnowledgeBaseORM.scope == KnowledgeScope.WORKSPACE.value,
                    KnowledgeBaseORM.owner_user_id == user_id,
                ),
            )
        )
        row = result.scalar_one_or_none()
        if row is None:
            raise ValueError("知识库不存在或无权访问")
        return row

    async def get_knowledge_base_for_action(
        self,
        *,
        kb_id: str,
        workspace_id: str,
        user_id: str,
        role: Any,
        action: str,
    ) -> KnowledgeBaseORM:
        """读取知识库并校验指定动作权限。"""
        row = await self.get_knowledge_base(kb_id=kb_id, workspace_id=workspace_id, user_id=user_id)
        self._ensure_kb_operation_allowed(row, user_id=user_id, role=role, action=action)
        return row

    async def update_knowledge_base(
        self,
        *,
        kb_id: str,
        workspace_id: str,
        user_id: str,
        role: Any,
        name: str | None = None,
        description: str | None = None,
        status: KnowledgeBaseStatus | None = None,
    ) -> KnowledgeBase:
        row = await self.get_knowledge_base_for_action(
            kb_id=kb_id,
            workspace_id=workspace_id,
            user_id=user_id,
            role=role,
            action="update",
        )
        if name is not None:
            row.name = name
        if description is not None:
            row.description = description
        if status is not None:
            row.status = status.value
        row.updated_at = utc_now_naive()
        await self.session.commit()
        await self.session.refresh(row)
        return _row_to_knowledge_base(row)

    async def delete_knowledge_base(
        self,
        *,
        kb_id: str,
        workspace_id: str,
        user_id: str,
        role: Any,
    ) -> None:
        row = await self.get_knowledge_base_for_action(
            kb_id=kb_id,
            workspace_id=workspace_id,
            user_id=user_id,
            role=role,
            action="delete",
        )
        documents = await self.session.execute(
            select(KnowledgeDocumentORM.storage_uri).where(KnowledgeDocumentORM.kb_id == row.id)
        )
        storage_paths = [path for path in documents.scalars().all() if path]
        await self.session.delete(row)
        await self.session.commit()
        for storage_path in storage_paths:
            Path(storage_path).unlink(missing_ok=True)

    async def add_document(
        self,
        *,
        kb_id: str,
        workspace_id: str,
        user_id: str,
        role: Any,
        file_name: str,
        content: bytes,
        metadata: dict[str, Any] | None = None,
        index_immediately: bool = True,
    ) -> KnowledgeDocument:
        kb = await self.get_knowledge_base_for_action(
            kb_id=kb_id,
            workspace_id=workspace_id,
            user_id=user_id,
            role=role,
            action="update",
        )
        file_type = validate_upload_file(file_name, content)
        await self._ensure_document_limit(kb_id=kb.id, file_name=file_name)

        document_id = str(uuid.uuid4())
        storage_uri = self._write_original_file(workspace_id, document_id, file_name, content)
        version_result = await self.session.execute(
            select(KnowledgeDocumentORM.version)
            .where(
                KnowledgeDocumentORM.kb_id == kb.id,
                KnowledgeDocumentORM.file_name == file_name,
            )
            .order_by(KnowledgeDocumentORM.version.desc())
            .limit(1)
        )
        latest_version = version_result.scalar_one_or_none() or 0

        document = KnowledgeDocumentORM(
            id=document_id,
            kb_id=kb.id,
            workspace_id=workspace_id,
            owner_user_id=kb.owner_user_id,
            workflow_run_id=kb.workflow_run_id,
            file_name=file_name,
            file_type=file_type,
            storage_uri=storage_uri,
            checksum=_hash_content(content),
            version=latest_version + 1,
            status=KnowledgeDocumentLifecycleStatus.ACTIVE.value,
            parse_status=KnowledgeDocumentStatus.PENDING.value,
            index_status=KnowledgeDocumentStatus.PENDING.value,
            metadata_json=metadata or {},
            created_by=user_id,
        )
        self.session.add(document)
        await self.session.flush()

        if index_immediately:
            await self._index_document_row(document, content, metadata or {})

        document.updated_at = utc_now_naive()
        kb.updated_at = utc_now_naive()
        await self.session.commit()
        await self.session.refresh(document)
        return _row_to_document(document)

    async def index_document(self, *, document_id: str) -> KnowledgeDocument:
        """执行单个文档的解析、切分和向量索引，供后台 worker 或同步路径复用。"""
        result = await self.session.execute(
            select(KnowledgeDocumentORM)
            .options(selectinload(KnowledgeDocumentORM.knowledge_base))
            .where(KnowledgeDocumentORM.id == document_id)
        )
        document = result.scalar_one_or_none()
        if document is None:
            raise ValueError("文档不存在")
        if not document.storage_uri:
            document.parse_status = KnowledgeDocumentStatus.FAILED.value
            document.index_status = KnowledgeDocumentStatus.FAILED.value
            document.error_message = "文档缺少原始文件，无法建立索引"
            document.updated_at = utc_now_naive()
            await self.session.commit()
            await self.session.refresh(document)
            return _row_to_document(document)

        content = Path(document.storage_uri).read_bytes()
        _validate_upload_content(document.file_type, content)
        await self._index_document_row(document, content, document.metadata_json or {})
        document.updated_at = utc_now_naive()
        if document.knowledge_base is not None:
            document.knowledge_base.updated_at = utc_now_naive()
        await self.session.commit()
        await self.session.refresh(document)
        return _row_to_document(document)

    async def _index_document_row(
        self,
        document: KnowledgeDocumentORM,
        content: bytes,
        metadata: dict[str, Any],
    ) -> None:
        """在当前事务中刷新一个文档的索引状态和索引行。"""
        document.parse_status = KnowledgeDocumentStatus.PROCESSING.value
        document.index_status = KnowledgeDocumentStatus.PROCESSING.value
        document.error_message = None
        await self.session.flush()
        try:
            text = self.parser.parse(document.file_name, content)
            await self._replace_document_chunks(document, text, metadata)
            await self.session.execute(
                update(KnowledgeDocumentORM)
                .where(
                    KnowledgeDocumentORM.kb_id == document.kb_id,
                    KnowledgeDocumentORM.file_name == document.file_name,
                    KnowledgeDocumentORM.id != document.id,
                    KnowledgeDocumentORM.version < document.version,
                    KnowledgeDocumentORM.status == KnowledgeDocumentLifecycleStatus.ACTIVE.value,
                )
                .values(
                    status=KnowledgeDocumentLifecycleStatus.ARCHIVED.value,
                    updated_at=utc_now_naive(),
                )
            )
            document.parse_status = KnowledgeDocumentStatus.READY.value
            document.index_status = KnowledgeDocumentStatus.READY.value
            document.error_message = None
        except Exception as exc:
            logger.warning("知识库文档索引失败 document_id=%s: %s", document.id, exc)
            document.parse_status = KnowledgeDocumentStatus.FAILED.value
            document.index_status = KnowledgeDocumentStatus.FAILED.value
            document.error_message = str(exc)

    async def _ensure_document_limit(self, *, kb_id: str, file_name: str) -> None:
        max_documents = get_settings().KNOWLEDGE_MAX_DOCUMENTS_PER_KB
        if max_documents <= 0:
            return
        result = await self.session.execute(
            select(KnowledgeDocumentORM.file_name)
            .where(
                KnowledgeDocumentORM.kb_id == kb_id,
                KnowledgeDocumentORM.status.in_(
                    [
                        KnowledgeDocumentLifecycleStatus.ACTIVE.value,
                        KnowledgeDocumentLifecycleStatus.DISABLED.value,
                    ]
                ),
            )
            .distinct()
        )
        active_file_names = set(result.scalars().all())
        if file_name in active_file_names:
            return
        if len(active_file_names) >= max_documents:
            raise ValueError(f"知识库文档数量不能超过 {max_documents} 个")

    def _write_original_file(
        self,
        workspace_id: str,
        document_id: str,
        file_name: str,
        content: bytes,
    ) -> str:
        storage_root = Path(get_settings().KNOWLEDGE_STORAGE_DIR)
        workspace_dir = storage_root / workspace_id
        workspace_dir.mkdir(parents=True, exist_ok=True)
        safe_suffix = Path(file_name).suffix or ".txt"
        target = workspace_dir / f"{document_id}{safe_suffix}"
        target.write_bytes(content)
        return str(target)

    async def _replace_document_chunks(
        self,
        document: KnowledgeDocumentORM,
        text: str,
        metadata: dict[str, Any],
    ) -> None:
        await self.session.execute(
            delete(KnowledgeChunkORM).where(KnowledgeChunkORM.document_id == document.id)
        )
        chunk_specs = self.chunker.split(text)
        embeddings = await self.embedding_provider.embed_documents(
            [spec["content"] for spec in chunk_specs]
        )
        for index, (spec, vector) in enumerate(zip(chunk_specs, embeddings, strict=False)):
            content = spec["content"]
            pgvector_value = vector if len(vector) == 1536 else None
            chunk = KnowledgeChunkORM(
                id=str(uuid.uuid4()),
                document_id=document.id,
                kb_id=document.kb_id,
                workspace_id=document.workspace_id,
                owner_user_id=document.owner_user_id,
                workflow_run_id=document.workflow_run_id,
                chunk_index=index,
                content=content,
                token_count=spec.get("token_count") or estimate_tokens(content),
                heading_path=spec.get("heading_path") or [],
                page_number=spec.get("page_number"),
                metadata_json={**metadata, **(spec.get("metadata") or {})},
                content_hash=_hash_content(content),
            )
            self.session.add(chunk)
            await self.session.flush()
            self.session.add(
                KnowledgeEmbeddingORM(
                    id=str(uuid.uuid4()),
                    chunk_id=chunk.id,
                    embedding_model=self.embedding_provider.effective_model_name,
                    vector_json=vector,
                    # pgvector 生产索引列固定为 1536 维；测试或本地 hash 维度不一致时仅保留 JSON 回退向量。
                    embedding_vector=pgvector_value,
                    dimension=len(vector),
                )
            )

    async def list_documents(
        self,
        *,
        kb_id: str,
        workspace_id: str,
        user_id: str,
    ) -> list[KnowledgeDocument]:
        await self.get_knowledge_base(kb_id=kb_id, workspace_id=workspace_id, user_id=user_id)
        result = await self.session.execute(
            select(KnowledgeDocumentORM)
            .where(KnowledgeDocumentORM.kb_id == kb_id)
            .order_by(
                KnowledgeDocumentORM.file_name.asc(),
                KnowledgeDocumentORM.version.desc(),
                KnowledgeDocumentORM.created_at.desc(),
            )
        )
        return [_row_to_document(row) for row in result.scalars().all()]

    async def get_document(
        self,
        *,
        document_id: str,
        workspace_id: str,
        user_id: str,
        role: Any | None = None,
        action: str = "read",
    ) -> KnowledgeDocumentORM:
        result = await self.session.execute(
            select(KnowledgeDocumentORM)
            .options(selectinload(KnowledgeDocumentORM.knowledge_base))
            .join(KnowledgeBaseORM, KnowledgeDocumentORM.kb_id == KnowledgeBaseORM.id)
            .where(
                KnowledgeDocumentORM.id == document_id,
                KnowledgeDocumentORM.workspace_id == workspace_id,
                or_(
                    KnowledgeBaseORM.scope == KnowledgeScope.WORKSPACE.value,
                    KnowledgeDocumentORM.owner_user_id == user_id,
                ),
            )
        )
        row = result.scalar_one_or_none()
        if row is None:
            raise ValueError("文档不存在或无权访问")
        if role is not None:
            self._ensure_kb_operation_allowed(
                row.knowledge_base,
                user_id=user_id,
                role=role,
                action=action,
            )
        return row

    async def delete_document(
        self,
        *,
        document_id: str,
        workspace_id: str,
        user_id: str,
        role: Any,
    ) -> None:
        document = await self.get_document(
            document_id=document_id,
            workspace_id=workspace_id,
            user_id=user_id,
            role=role,
            action="delete",
        )
        storage_uri = document.storage_uri
        await self.session.delete(document)
        await self.session.commit()
        if storage_uri:
            Path(storage_uri).unlink(missing_ok=True)

    async def update_document_status(
        self,
        *,
        document_id: str,
        workspace_id: str,
        user_id: str,
        role: Any,
        status: KnowledgeDocumentLifecycleStatus,
    ) -> KnowledgeDocument:
        """更新文档是否参与检索，支持停用和归档旧资料。"""
        document = await self.get_document(
            document_id=document_id,
            workspace_id=workspace_id,
            user_id=user_id,
            role=role,
            action="update",
        )
        document.status = status.value
        document.updated_at = utc_now_naive()
        await self.session.commit()
        await self.session.refresh(document)
        return _row_to_document(document)

    async def reindex_document(
        self,
        *,
        document_id: str,
        workspace_id: str,
        user_id: str,
        role: Any,
        index_immediately: bool = True,
    ) -> KnowledgeDocument:
        document = await self.get_document(
            document_id=document_id,
            workspace_id=workspace_id,
            user_id=user_id,
            role=role,
            action="update",
        )
        if not document.storage_uri:
            raise ValueError("文档缺少原始文件，无法重建索引")
        if index_immediately:
            await self._index_document_row(document, Path(document.storage_uri).read_bytes(), {})
        else:
            document.parse_status = KnowledgeDocumentStatus.PENDING.value
            document.index_status = KnowledgeDocumentStatus.PENDING.value
            document.error_message = None
        document.updated_at = utc_now_naive()
        await self.session.commit()
        await self.session.refresh(document)
        return _row_to_document(document)

    async def search(
        self,
        *,
        request: KnowledgeSearchRequest,
        workspace_id: str,
        user_id: str,
    ) -> KnowledgeSearchResponse:
        query = request.query.strip()
        queries = self._rewrite_queries(query, request)
        query_vector = await self.embedding_provider.embed_query(" ".join(queries))
        rows = await self._load_search_rows(
            workspace_id=workspace_id,
            user_id=user_id,
            scopes=request.scopes,
            filters=request.filters,
            workflow_run_id=request.workflow_run_id,
            query_vector=query_vector,
            mode=request.mode,
            candidate_limit=max(request.top_k * 5, 50),
        )
        candidates = [
            self._score_candidate(
                row,
                query=query,
                queries=queries,
                query_vector=query_vector,
                mode=request.mode,
            )
            for row in rows
        ]
        if request.enable_rerank:
            candidates = self.reranker_provider.rerank(query, candidates)
        candidates = [candidate for candidate in candidates if candidate.score >= request.min_score]
        candidates = sorted(candidates, key=lambda item: item.score, reverse=True)[: request.top_k]
        evidence_chunks = [
            self._candidate_to_evidence_chunk(candidate, query, request.enable_context_compression)
            for candidate in candidates
        ]
        conflicts = (
            self._detect_conflicts(evidence_chunks) if request.enable_conflict_detection else []
        )
        evidence_pack = EvidencePack(
            query=query,
            rewritten_queries=queries[1:],
            scopes=request.scopes,
            chunks=evidence_chunks,
            conflicts=conflicts,
            unverified_points=[] if evidence_chunks else [query],
            retrieval_mode=request.mode,
        )
        retrieval_log_id = await self._create_retrieval_log(
            request=request,
            workspace_id=workspace_id,
            user_id=user_id,
            queries=queries,
            evidence_pack=evidence_pack,
        )
        return KnowledgeSearchResponse(
            evidence_pack=evidence_pack,
            retrieval_log_id=retrieval_log_id,
        )

    async def evaluate_retrieval(
        self,
        *,
        request: RetrievalEvaluationRequest,
        workspace_id: str,
        user_id: str,
    ) -> RetrievalEvaluationResponse:
        """对一组查询执行检索质量评测，返回 hit rate / MRR / Precision@k。"""
        results: list[RetrievalEvaluationResult] = []
        empty_expected_count = 0

        for case in request.cases:
            expected_document_ids = set(case.expected_document_ids)
            expected_chunk_ids = set(case.expected_chunk_ids)
            if not expected_document_ids and not expected_chunk_ids:
                empty_expected_count += 1

            search_response = await self.search(
                request=KnowledgeSearchRequest(
                    query=case.query,
                    scopes=request.scopes,
                    top_k=request.top_k,
                    min_score=request.min_score,
                    mode=request.mode,
                    filters=request.filters,
                    enable_query_rewrite=request.enable_query_rewrite,
                    enable_multi_query=request.enable_multi_query,
                    enable_rerank=request.enable_rerank,
                    enable_context_compression=request.enable_context_compression,
                    enable_conflict_detection=request.enable_conflict_detection,
                ),
                workspace_id=workspace_id,
                user_id=user_id,
            )
            chunks = search_response.evidence_pack.chunks
            retrieved_document_ids = [chunk.document_id for chunk in chunks]
            retrieved_chunk_ids = [chunk.chunk_id for chunk in chunks]
            relevant_ranks = [
                index
                for index, chunk in enumerate(chunks, start=1)
                if chunk.document_id in expected_document_ids
                or chunk.chunk_id in expected_chunk_ids
            ]
            first_relevant_rank = min(relevant_ranks) if relevant_ranks else None
            relevant_count = len(relevant_ranks)
            precision_at_k = relevant_count / max(1, len(chunks))
            results.append(
                RetrievalEvaluationResult(
                    case_id=case.id,
                    query=case.query,
                    expected_document_ids=list(expected_document_ids),
                    expected_chunk_ids=list(expected_chunk_ids),
                    retrieved_document_ids=retrieved_document_ids,
                    retrieved_chunk_ids=retrieved_chunk_ids,
                    hit=first_relevant_rank is not None,
                    first_relevant_rank=first_relevant_rank,
                    reciprocal_rank=round(1 / first_relevant_rank, 4)
                    if first_relevant_rank
                    else 0.0,
                    precision_at_k=round(precision_at_k, 4),
                )
            )

        total_cases = len(results)
        hit_count = sum(1 for result in results if result.hit)
        return RetrievalEvaluationResponse(
            summary=RetrievalEvaluationSummary(
                total_cases=total_cases,
                hit_count=hit_count,
                hit_rate=round(hit_count / max(1, total_cases), 4),
                mean_reciprocal_rank=round(
                    sum(result.reciprocal_rank for result in results) / max(1, total_cases),
                    4,
                ),
                mean_precision_at_k=round(
                    sum(result.precision_at_k for result in results) / max(1, total_cases),
                    4,
                ),
                empty_expected_count=empty_expected_count,
            ),
            results=results,
        )

    async def get_usage_stats(
        self,
        *,
        workspace_id: str,
        user_id: str,
        role: Any | None = None,
    ) -> KnowledgeUsageStats:
        """基于检索日志统计当前用户可见范围内的知识库使用情况。"""
        filters = [KnowledgeRetrievalLogORM.workspace_id == workspace_id]
        if role is None or not check_permission(
            _normalize_role_value(role),
            "knowledge_base",
            "manage",
        ):
            filters.append(KnowledgeRetrievalLogORM.user_id == user_id)

        result = await self.session.execute(select(KnowledgeRetrievalLogORM).where(and_(*filters)))
        logs = list(result.scalars().all())
        total_searches = len(logs)
        total_chunks = 0
        conflict_search_count = 0
        unverified_search_count = 0
        scope_counts: Counter[str] = Counter()
        mode_counts: Counter[str] = Counter()
        last_search_at = None

        for log in logs:
            retrieved_chunk_ids = log.retrieved_chunk_ids or []
            if isinstance(retrieved_chunk_ids, list):
                total_chunks += len(retrieved_chunk_ids)
            if log.created_at and (last_search_at is None or log.created_at > last_search_at):
                last_search_at = log.created_at

            for scope in log.retrieval_scope or []:
                scope_counts[str(scope)] += 1
            mode_counts[str(log.retrieval_mode)] += 1

            metadata = log.metadata_json or {}
            if metadata.get("conflicts"):
                conflict_search_count += 1
            if metadata.get("unverified_points"):
                unverified_search_count += 1

        return KnowledgeUsageStats(
            workspace_id=workspace_id,
            total_searches=total_searches,
            total_chunks_returned=total_chunks,
            average_chunks_per_search=round(total_chunks / max(1, total_searches), 2),
            conflict_search_count=conflict_search_count,
            unverified_search_count=unverified_search_count,
            last_search_at=last_search_at,
            scope_counts=dict(scope_counts),
            mode_counts=dict(mode_counts),
        )

    def _rewrite_queries(self, query: str, request: KnowledgeSearchRequest) -> list[str]:
        queries = [query]
        if request.enable_query_rewrite:
            normalized = re.sub(r"\s+", " ", query).strip()
            if normalized and normalized != query:
                queries.append(normalized)
        if request.enable_multi_query:
            queries.append(f"{query} 依据 来源 引用")
            queries.append(f"{query} 风险 冲突 事实")
        return list(dict.fromkeys(item for item in queries if item.strip()))

    async def _load_search_rows(
        self,
        *,
        workspace_id: str,
        user_id: str,
        scopes: list[KnowledgeScope],
        filters: dict[str, Any],
        workflow_run_id: str | None,
        query_vector: list[float],
        mode: RetrievalMode,
        candidate_limit: int,
    ) -> list[
        tuple[
            KnowledgeChunkORM, KnowledgeDocumentORM, KnowledgeBaseORM, KnowledgeEmbeddingORM | None
        ]
    ]:
        scope_values = [scope.value for scope in scopes]
        conditions = [
            KnowledgeChunkORM.workspace_id == workspace_id,
            KnowledgeBaseORM.status == KnowledgeBaseStatus.ACTIVE.value,
            KnowledgeDocumentORM.status == KnowledgeDocumentLifecycleStatus.ACTIVE.value,
            KnowledgeDocumentORM.index_status == KnowledgeDocumentStatus.READY.value,
            KnowledgeBaseORM.scope.in_(scope_values),
            or_(
                KnowledgeBaseORM.scope == KnowledgeScope.WORKSPACE.value,
                KnowledgeChunkORM.owner_user_id == user_id,
            ),
        ]
        file_types = filters.get("file_type")
        if isinstance(file_types, list) and file_types:
            conditions.append(KnowledgeDocumentORM.file_type.in_(file_types))
        kb_ids = filters.get("kb_ids")
        if isinstance(kb_ids, list) and kb_ids:
            conditions.append(KnowledgeBaseORM.id.in_(kb_ids))
        if KnowledgeScope.RUN_UPLOAD in scopes:
            if workflow_run_id:
                # 本次运行上传资料必须绑定当前运行，避免跨运行串用临时证据。
                conditions.append(
                    or_(
                        KnowledgeBaseORM.scope != KnowledgeScope.RUN_UPLOAD.value,
                        KnowledgeBaseORM.workflow_run_id == workflow_run_id,
                    )
                )
            else:
                conditions.append(KnowledgeBaseORM.scope != KnowledgeScope.RUN_UPLOAD.value)

        if self._should_use_pgvector_candidates(query_vector=query_vector, mode=mode):
            result = await self.session.execute(
                select(
                    KnowledgeChunkORM,
                    KnowledgeDocumentORM,
                    KnowledgeBaseORM,
                    KnowledgeEmbeddingORM,
                )
                .join(
                    KnowledgeDocumentORM, KnowledgeChunkORM.document_id == KnowledgeDocumentORM.id
                )
                .join(KnowledgeBaseORM, KnowledgeChunkORM.kb_id == KnowledgeBaseORM.id)
                .join(KnowledgeEmbeddingORM, KnowledgeEmbeddingORM.chunk_id == KnowledgeChunkORM.id)
                .where(
                    and_(
                        *conditions,
                        KnowledgeEmbeddingORM.embedding_vector.is_not(None),
                    )
                )
                .order_by(KnowledgeEmbeddingORM.embedding_vector.cosine_distance(query_vector))
                .limit(candidate_limit)
            )
            return list(result.all())

        result = await self.session.execute(
            select(KnowledgeChunkORM, KnowledgeDocumentORM, KnowledgeBaseORM, KnowledgeEmbeddingORM)
            .join(KnowledgeDocumentORM, KnowledgeChunkORM.document_id == KnowledgeDocumentORM.id)
            .join(KnowledgeBaseORM, KnowledgeChunkORM.kb_id == KnowledgeBaseORM.id)
            .outerjoin(
                KnowledgeEmbeddingORM, KnowledgeEmbeddingORM.chunk_id == KnowledgeChunkORM.id
            )
            .where(and_(*conditions))
        )
        return list(result.all())

    def _should_use_pgvector_candidates(
        self,
        *,
        query_vector: list[float],
        mode: RetrievalMode,
    ) -> bool:
        if mode == RetrievalMode.KEYWORD:
            return False
        if len(query_vector) != 1536:
            return False
        if not hasattr(KnowledgeEmbeddingORM, "embedding_vector"):
            return False
        bind = self.session.get_bind()
        return bool(bind is not None and bind.dialect.name == "postgresql")

    def _score_candidate(
        self,
        row: tuple[
            KnowledgeChunkORM, KnowledgeDocumentORM, KnowledgeBaseORM, KnowledgeEmbeddingORM | None
        ],
        *,
        query: str,
        queries: list[str],
        query_vector: list[float],
        mode: RetrievalMode,
    ) -> _Candidate:
        chunk, document, kb, embedding = row
        vector_score = (
            _cosine_similarity(query_vector, list(embedding.vector_json or []))
            if embedding
            else 0.0
        )
        keyword_score = max(_keyword_score(item, chunk.content) for item in queries)
        if mode == RetrievalMode.VECTOR:
            score = vector_score
        elif mode == RetrievalMode.KEYWORD:
            score = keyword_score
        else:
            score = vector_score * 0.7 + keyword_score * 0.3
        return _Candidate(
            chunk=chunk,
            document=document,
            kb=kb,
            embedding=embedding,
            score=round(score, 4),
            vector_score=round(vector_score, 4),
            keyword_score=round(keyword_score, 4),
        )

    def _candidate_to_evidence_chunk(
        self,
        candidate: _Candidate,
        query: str,
        compress: bool,
    ) -> EvidenceChunk:
        content = candidate.chunk.content
        if compress:
            content = self._compress_content(query, content)
        return EvidenceChunk(
            chunk_id=candidate.chunk.id,
            document_id=candidate.document.id,
            knowledge_base_id=candidate.kb.id,
            document_name=candidate.document.file_name,
            scope=KnowledgeScope(candidate.kb.scope),
            page_number=candidate.chunk.page_number,
            heading_path=candidate.chunk.heading_path or [],
            score=candidate.score,
            vector_score=candidate.vector_score,
            keyword_score=candidate.keyword_score,
            rerank_score=candidate.rerank_score,
            content=content,
            metadata=candidate.chunk.metadata_json or {},
        )

    def _compress_content(self, query: str, content: str) -> str:
        if len(content) <= 1600:
            return content
        terms = _tokenize(query)
        lowered = content.lower()
        first_match = min(
            (lowered.find(term) for term in terms if lowered.find(term) >= 0),
            default=0,
        )
        start = max(0, first_match - 500)
        end = min(len(content), start + 1600)
        prefix = "..." if start > 0 else ""
        suffix = "..." if end < len(content) else ""
        return f"{prefix}{content[start:end]}{suffix}"

    def _detect_conflicts(self, chunks: list[EvidenceChunk]) -> list[KnowledgeConflict]:
        conflicts: list[KnowledgeConflict] = []
        negative_markers = {"不是", "不能", "不得", "不应", "禁止", "无", "not", "never"}
        for index, left in enumerate(chunks):
            left_terms = set(_tokenize(left.content))
            for right in chunks[index + 1 :]:
                overlap = left_terms & set(_tokenize(right.content))
                if len(overlap) < 4:
                    continue
                left_has_negative = any(
                    marker in left.content.lower() for marker in negative_markers
                )
                right_has_negative = any(
                    marker in right.content.lower() for marker in negative_markers
                )
                if left_has_negative != right_has_negative:
                    conflicts.append(
                        KnowledgeConflict(
                            topic=" / ".join(list(overlap)[:4]),
                            chunk_ids=[left.chunk_id, right.chunk_id],
                            reason="相近主题证据存在肯定/否定表述差异，请人工复核。",
                        )
                    )
        return conflicts

    async def _create_retrieval_log(
        self,
        *,
        request: KnowledgeSearchRequest,
        workspace_id: str,
        user_id: str,
        queries: list[str],
        evidence_pack: EvidencePack,
    ) -> str:
        log_id = str(uuid.uuid4())
        self.session.add(
            KnowledgeRetrievalLogORM(
                id=log_id,
                workflow_run_id=request.workflow_run_id,
                node_run_id=request.node_run_id,
                user_id=user_id,
                workspace_id=workspace_id,
                query=request.query,
                rewritten_queries=queries[1:],
                retrieval_scope=[scope.value for scope in request.scopes],
                retrieval_mode=request.mode.value,
                top_k=request.top_k,
                min_score=request.min_score,
                retrieved_chunk_ids=[chunk.chunk_id for chunk in evidence_pack.chunks],
                scores=[
                    {
                        "chunk_id": chunk.chunk_id,
                        "score": chunk.score,
                        "vector_score": chunk.vector_score,
                        "keyword_score": chunk.keyword_score,
                        "rerank_score": chunk.rerank_score,
                    }
                    for chunk in evidence_pack.chunks
                ],
                metadata_json={
                    "conflicts": [conflict.model_dump() for conflict in evidence_pack.conflicts],
                    "unverified_points": evidence_pack.unverified_points,
                },
            )
        )
        await self.session.commit()
        return log_id
