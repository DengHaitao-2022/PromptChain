"""工作流最终产物导出服务。"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Any


@dataclass(slots=True)
class ExportSection:
    """导出章节的统一结构。"""

    id: str
    title: str
    content: str
    word_count: int = 0


@dataclass(slots=True)
class WorkflowExportPayload:
    """导出时使用的最终产物载荷。"""

    title: str
    abstract: str
    sections: list[ExportSection]
    citations: list[dict[str, Any]]
    raw: dict[str, Any] | list[Any] | str | None


def _is_record(value: Any) -> bool:
    return isinstance(value, dict)


def _as_string_array(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, str)]


def _count_words(content: str) -> int:
    compact = re.sub(r"\s+", "", content)
    return len(compact) or len(content)


def _format_entry_label(key: str) -> str:
    return " ".join(part[:1].upper() + part[1:] for part in key.split("_") if part)


def _parse_generated_content(content: Any) -> list[ExportSection]:
    if not isinstance(content, str) or not content.strip():
        return []

    sections = re.split(r"(?:\n## |^## )", content, flags=re.MULTILINE)
    if len(sections) <= 1:
        body = content.strip()
        return [
            ExportSection(
                id="generated_body",
                title="正文草稿",
                content=body,
                word_count=_count_words(body),
            )
        ]

    result: list[ExportSection] = []
    for index, section in enumerate(sections):
        if not section.strip():
            continue
        lines = section.splitlines()
        title = lines[0].strip() or f"章节 {index + 1}"
        body = "\n".join(lines[1:]).strip()
        result.append(
            ExportSection(
                id=f"gen_sec_{index}",
                title=title,
                content=body,
                word_count=_count_words(body),
            )
        )
    return result


def _get_outline_title_map(outline: Any) -> dict[str, str]:
    if not _is_record(outline) or not isinstance(outline.get("sections"), list):
        return {}

    title_map: dict[str, str] = {}

    def visit(sections: list[Any]) -> None:
        for section in sections:
            if not _is_record(section):
                continue
            section_id = section.get("id")
            title = section.get("title")
            if isinstance(section_id, str) and isinstance(title, str):
                title_map[section_id] = title
            if isinstance(section.get("subsections"), list):
                visit(section["subsections"])

    visit(outline["sections"])
    return title_map


def _get_content_text(value: Any) -> str:
    if isinstance(value, str):
        return value
    if not _is_record(value):
        return ""

    for key in ("compiled_content", "content", "text", "markdown"):
        item = value.get(key)
        if isinstance(item, str):
            return item

    sections = value.get("sections")
    if _is_record(sections):
        blocks: list[str] = []
        for section_id, section_content in sections.items():
            if isinstance(section_content, str):
                text = section_content
            else:
                text = json.dumps(section_content, ensure_ascii=False, indent=2)
            blocks.append(f"## {_format_entry_label(str(section_id))}\n{text}")
        return "\n\n".join(blocks)

    return ""


def _get_artifact_timestamp(artifact: dict[str, Any]) -> float:
    created_at = artifact.get("created_at")
    version = artifact.get("version")
    if isinstance(created_at, str):
        try:
            return datetime.fromisoformat(created_at.replace("Z", "+00:00")).timestamp()
        except ValueError:
            return 0.0
    if isinstance(version, (int, float)):
        return float(version)
    return 0.0


def _get_latest_artifact_by_type(
    trace: dict[str, Any] | None,
    artifact_type: str,
    preferred_id: Any = None,
) -> dict[str, Any] | None:
    if not _is_record(trace):
        return None

    artifacts = trace.get("artifacts")
    if not _is_record(artifacts):
        return None

    if isinstance(preferred_id, str):
        preferred = artifacts.get(preferred_id)
        if _is_record(preferred) and preferred.get("type") == artifact_type:
            return preferred

    candidates = [
        artifact
        for artifact in artifacts.values()
        if _is_record(artifact) and artifact.get("type") == artifact_type
    ]
    if not candidates:
        return None

    return sorted(candidates, key=_get_artifact_timestamp, reverse=True)[0]


def normalize_final_content(final_content: Any) -> list[ExportSection]:
    if not _is_record(final_content):
        return []

    sections: list[ExportSection] = []
    for key, value in final_content.items():
        title = _format_entry_label(str(key))

        if _is_record(value):
            content = value.get("content") or value.get("preview") or value.get("text") or ""
            raw_word_count = value.get("word_count") or value.get("wordCount") or 0
        else:
            content = value
            raw_word_count = 0

        text = str(content or "").strip()
        if not text:
            continue

        try:
            word_count = int(raw_word_count) if isinstance(raw_word_count, (int, float, str)) else 0
        except (TypeError, ValueError):
            word_count = 0
        sections.append(
            ExportSection(
                id=str(key),
                title=title,
                content=text,
                word_count=word_count or _count_words(text),
            )
        )

    return sections


def _build_sections_from_final_artifact(
    artifact: dict[str, Any] | None,
    outline: Any,
) -> list[ExportSection]:
    if not _is_record(artifact):
        return []

    content = artifact.get("content")
    title_map = _get_outline_title_map(outline)

    if _is_record(content) and _is_record(content.get("sections")):
        sections = content["sections"]
        ordered_ids = _as_string_array(content.get("section_order"))
        section_ids = ordered_ids or list(sections.keys())
        result: list[ExportSection] = []

        for section_id in section_ids:
            section_content = sections.get(section_id)
            text = (
                section_content
                if isinstance(section_content, str)
                else json.dumps(section_content, ensure_ascii=False, indent=2)
            )
            text = text.strip()
            if not text:
                continue
            result.append(
                ExportSection(
                    id=section_id,
                    title=title_map.get(section_id, _format_entry_label(section_id)),
                    content=text,
                    word_count=_count_words(text),
                )
            )

        return result

    sections = _parse_generated_content(_get_content_text(content))
    if sections:
        return sections

    return normalize_final_content(content)


def _build_sections_from_section_artifacts(
    trace: dict[str, Any] | None,
    outline: Any,
) -> list[ExportSection]:
    if not _is_record(trace):
        return []

    artifacts = trace.get("artifacts")
    if not _is_record(artifacts):
        return []

    title_map = _get_outline_title_map(outline)
    latest_by_section: dict[str, tuple[ExportSection, float]] = {}

    for artifact in artifacts.values():
        if not _is_record(artifact) or artifact.get("type") != "section_content":
            continue

        content = artifact.get("content")
        if not _is_record(content):
            continue

        section_id = content.get("section_id") or artifact.get("id") or ""
        text = content.get("content") if isinstance(content.get("content"), str) else ""
        if not isinstance(section_id, str) or not text.strip():
            continue

        timestamp = _get_artifact_timestamp(artifact)
        current = latest_by_section.get(section_id)
        if current and current[1] >= timestamp:
            continue

        latest_by_section[section_id] = (
            ExportSection(
                id=section_id,
                title=(
                    content.get("section_title")
                    if isinstance(content.get("section_title"), str)
                    else title_map.get(section_id, _format_entry_label(section_id))
                ),
                content=text.strip(),
                word_count=_count_words(text),
            ),
            timestamp,
        )

    ordered_ids = list(title_map.keys())
    sections = [item[0] for item in latest_by_section.values()]

    def sort_key(section: ExportSection) -> tuple[int, str]:
        if section.id in ordered_ids:
            return (ordered_ids.index(section.id), section.id)
        return (len(ordered_ids) + 1, section.id)

    return sorted(sections, key=sort_key)


def _extract_citations(raw: Any, graph_state: dict[str, Any]) -> list[dict[str, Any]]:
    """从最终产物或图状态中提取结构化引用来源。"""
    candidates: Any = None
    if _is_record(raw):
        candidates = raw.get("citations")
    if not isinstance(candidates, list):
        candidates = graph_state.get("citations")
    if not isinstance(candidates, list):
        return []

    citations: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for item in candidates:
        if not _is_record(item):
            continue
        chunk_id = item.get("chunk_id")
        document_id = item.get("document_id")
        document_name = item.get("document_name")
        if not isinstance(document_name, str) or not document_name.strip():
            continue
        key = (str(chunk_id or ""), str(document_id or document_name))
        if key in seen:
            continue
        seen.add(key)
        citations.append(
            {
                "chunk_id": chunk_id if isinstance(chunk_id, str) else None,
                "document_id": document_id if isinstance(document_id, str) else None,
                "document_name": document_name.strip(),
                "scope": item.get("scope") if isinstance(item.get("scope"), str) else None,
                "score": item.get("score") if isinstance(item.get("score"), (int, float)) else None,
                "page_number": item.get("page_number")
                if isinstance(item.get("page_number"), int)
                else None,
                "heading_path": item.get("heading_path")
                if isinstance(item.get("heading_path"), list)
                else [],
            }
        )
    return citations


def build_workflow_export_payload(
    graph_state: dict[str, Any] | None,
    trace: dict[str, Any] | None,
    *,
    workflow_name: str | None = None,
) -> WorkflowExportPayload:
    """按前端完成态相同优先级构建最终导出载荷。"""
    state = graph_state if _is_record(graph_state) else {}
    outline = state.get("outline")
    trace_workflow = trace.get("workflow") if _is_record(trace) else {}
    final_artifact_id = state.get("final_content_artifact_id")
    if not isinstance(final_artifact_id, str) and _is_record(trace_workflow):
        value = trace_workflow.get("final_artifact_id")
        final_artifact_id = value if isinstance(value, str) else None

    final_artifact = _get_latest_artifact_by_type(trace, "final_content", final_artifact_id)
    sections = _build_sections_from_final_artifact(final_artifact, outline)
    if not sections:
        sections = _build_sections_from_section_artifacts(trace, outline)
    if not sections:
        sections = _parse_generated_content(state.get("generated_content"))
    if not sections:
        sections = normalize_final_content(state.get("final_content"))

    title = workflow_name or "生成内容"
    abstract = ""
    if _is_record(outline):
        if isinstance(outline.get("title"), str) and outline["title"].strip():
            title = outline["title"].strip()
        if isinstance(outline.get("abstract"), str):
            abstract = outline["abstract"].strip()

    raw = (
        final_artifact.get("content") if _is_record(final_artifact) else state.get("final_content")
    )
    return WorkflowExportPayload(
        title=title,
        abstract=abstract,
        sections=sections,
        citations=_extract_citations(raw, state),
        raw=raw,
    )


def _format_citation_label(citation: dict[str, Any], index: int) -> str:
    document_name = citation.get("document_name") or f"来源 {index}"
    details: list[str] = []
    scope = citation.get("scope")
    if isinstance(scope, str) and scope:
        details.append(scope)
    page_number = citation.get("page_number")
    if isinstance(page_number, int):
        details.append(f"第 {page_number} 页")
    heading_path = citation.get("heading_path")
    if isinstance(heading_path, list):
        heading = " / ".join(str(item) for item in heading_path if item)
        if heading:
            details.append(heading)
    score = citation.get("score")
    if isinstance(score, (int, float)):
        details.append(f"score={score:.3f}")

    suffix = f"（{'；'.join(details)}）" if details else ""
    return f"[{index}] {document_name}{suffix}"


def build_markdown(payload: WorkflowExportPayload) -> str:
    """将最终产物拼接成 Markdown 文档。"""
    blocks = [f"# {payload.title}"]

    if payload.abstract:
        blocks.append(f"## 摘要\n\n{payload.abstract}")

    for index, section in enumerate(payload.sections, start=1):
        blocks.append(f"## {index}. {section.title}\n\n{section.content.strip()}")

    if payload.citations:
        citation_lines = [
            _format_citation_label(citation, index)
            for index, citation in enumerate(payload.citations, start=1)
        ]
        blocks.append("## 参考来源\n\n" + "\n".join(f"- {line}" for line in citation_lines))

    return "\n\n".join(blocks).strip() + "\n"


def _flush_code_block(document: Any, code_lines: list[str]) -> None:
    if not code_lines:
        return

    from docx.shared import Pt

    paragraph = document.add_paragraph()
    run = paragraph.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def add_markdown_like_content(document: Any, content: str) -> None:
    """按最小规则把 Markdown 文本写入 DOCX。"""
    in_code_block = False
    code_lines: list[str] = []

    for raw_line in content.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code_block:
                _flush_code_block(document, code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if not stripped:
            continue

        if stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=3)
            continue
        if stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=2)
            continue
        if stripped.startswith("# "):
            document.add_heading(stripped[2:].strip(), level=1)
            continue
        if stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
            continue
        if re.match(r"^\d+\.\s+", stripped):
            document.add_paragraph(re.sub(r"^\d+\.\s+", "", stripped), style="List Number")
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(stripped.lstrip("> ").strip())
            run.italic = True
            continue

        document.add_paragraph(stripped)

    if in_code_block:
        _flush_code_block(document, code_lines)


def build_docx(payload: WorkflowExportPayload) -> BytesIO:
    """生成可下载的 DOCX 文件内容。"""
    from docx import Document

    document = Document()
    document.add_heading(payload.title, level=0)

    if payload.abstract:
        document.add_heading("摘要", level=1)
        add_markdown_like_content(document, payload.abstract)

    if payload.sections:
        document.add_heading("目录", level=1)
        for index, section in enumerate(payload.sections, start=1):
            document.add_paragraph(f"{index}. {section.title}", style="List Number")

    for index, section in enumerate(payload.sections, start=1):
        document.add_heading(f"{index}. {section.title}", level=1)
        add_markdown_like_content(document, section.content)

    if payload.citations:
        document.add_heading("参考来源", level=1)
        for index, citation in enumerate(payload.citations, start=1):
            document.add_paragraph(_format_citation_label(citation, index), style="List Bullet")

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
